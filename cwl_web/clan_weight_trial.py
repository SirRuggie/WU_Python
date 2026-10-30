"""Staff-only official weight editor using live membership and durable player weights."""
from datetime import datetime, timezone
from io import BytesIO
from aiohttp import web
from .providers import tag, ProviderUnavailable
from utils.clan_weight_store import WeightConflict
from .clan_weights import observations


def authorize(request):
    from .app import _require, _owner, _store, _fail
    from .access_control import request_scopes
    actor = _require(request)
    if not _owner(actor) and 'clan_weights' not in request_scopes(request, actor['user_id']):
        _fail(403, 'Clan management access required')
    return actor


async def family(request):
    result = await request.app['providers'].family_clans()
    return [clan for clan in result['clans'] if clan.get('type') == 'FWA']


async def current(request):
    from .app import _fail
    try:
        value = tag(request.match_info['tag'])
    except ValueError:
        _fail(400, 'Invalid clan tag')
    clans = await family(request)
    if value not in {row['tag'] for row in clans}:
        _fail(404, 'This clan is not an FWA clan in the family. Refresh the clan list.')
    snapshot = await observations(request.app).refresh_roster(value, force=True)
    if snapshot['status'] != 'current':
        raise ProviderUnavailable('Current clan members could not be verified. Try again.')
    return snapshot['value']


def response(roster, saved):
    return {'clan': {'tag':roster['tag'],'name':roster['name']},
            'checked_at':datetime.now(timezone.utc).isoformat(),
            'members':[{**member,'weight':saved.get(member['tag'],{}).get('weight'),
                        'storages':saved.get(member['tag'],{}).get('storages') or (5 if (member.get('town_hall') or 18) > 8 else 4),
                        'revision':saved.get(member['tag'],{}).get('revision',0)} for member in roster['members']]}


def last_saved(request, clan_tag):
    saved = request.app['clan_weight_trial'].last_saved(clan_tag)
    if not saved:
        return None
    actor_id = saved['actor_id'].split(':', 1)[0]
    name = saved['actor_name']
    if not name:
        from .app import _store
        rows = _store(request).rows(
            "SELECT COALESCE(NULLIF(display_name,''),username) AS name FROM sessions WHERE user_id=? ORDER BY expires_at DESC LIMIT 1",
            (actor_id,),
        )
        name = rows[0]['name'] if rows else None
    return {'at': saved['updated_at'], 'discord_id': actor_id,
            'name': name or ('Discord user ' + actor_id if actor_id.isdigit() else 'Unknown user'),
            'imported': saved['actor_id'].endswith((':sheet-test-import', ':sheet-import'))}


async def directory(request):
    authorize(request)
    clans = await family(request)
    cache = observations(request.app)
    snapshots = {clan['tag']: cache.roster(clan['tag']) for clan in clans}
    member_tags = {member['tag'] for snapshot in snapshots.values()
                   for member in (snapshot.get('value') or {}).get('members', [])}
    ledger = request.app['clan_weight_trial']
    saved = ledger.for_members(member_tags)
    result = []
    for clan in clans:
        snapshot = snapshots[clan['tag']]
        roster = snapshot['value']
        count = len(roster['members']) if roster is not None else None
        weighed = sum(member['tag'] in saved for member in roster['members']) if roster is not None else None
        last = ledger.last_saved(clan['tag'])
        result.append({**clan, 'members_count': count, 'weighed_count': weighed,
                       'missing_count': count - weighed if count is not None else None,
                       'last_saved_at': last['updated_at'] if last else None,
                       'membership_checked_at': snapshot['checked_at'], 'status': snapshot['status']})
    result.sort(key=lambda clan: (not bool(clan['missing_count']),
                                 clan['name'].casefold(), clan['tag']))
    outside = None
    outside_status = 'unavailable'
    season_id = None
    if request.query.get('season_id'):
        from .app import _require, _store, _season, _roster_rows, _fail
        _require(request, 'editor')  # Same authorization as Build roster.
        try:
            season_id = int(request.query['season_id'])
        except ValueError:
            _fail(400, 'Choose a valid season.')
        with _store(request).tx() as db:
            _season(db, season_id)
            accounts = _roster_rows(db, season_id)
        if all(snapshot['status'] == 'current' for snapshot in snapshots.values()):
            outside = sum(account['effective_weight'] is None and not account['clan_weight_managed']
                          and account['tag'] not in member_tags for account in accounts)
            outside_status = 'current'
        else:
            outside_status = 'loading' if any(s['status'] == 'loading' for s in snapshots.values()) else 'stale'
    return web.json_response({'clans': result, 'season_id': season_id,
                              'outside_fwa_missing_count': outside,
                              'outside_fwa_status': outside_status})


def saved_player(row):
    return {'tag': row['tag'], 'name': row['player_name'], 'town_hall': row['town_hall'],
            'weight': row['weight'], 'storages': row['storages'], 'revision': row['revision'],
            'last_saved_clan_tag': row['last_clan_tag'], 'last_saved_clan_name': row['last_clan_name'],
            'updated_at': row['updated_at']}


async def saved_players(request):
    authorize(request)
    from .app import _fail
    query = request.query.get('q', '').strip()
    try:
        limit = int(request.query.get('limit', '25'))
        offset = int(request.query.get('offset', '0'))
    except ValueError:
        _fail(400, 'Use valid search page numbers.')
    if len(query) > 80 or not 1 <= limit <= 25 or not 0 <= offset <= 100000:
        _fail(400, 'Use a search under 80 characters and a page of up to 25 players.')
    found = request.app['clan_weight_trial'].search(query, limit=limit, offset=offset)
    try:
        directory = await request.app['providers'].family_clans()
        clans = {clan['tag']: clan for clan in directory['clans']}
    except ProviderUnavailable:
        clans = None
    cache = observations(request.app)

    async def locate(row):
        snapshot = await cache.player(row['tag'])
        current = snapshot['value'] if snapshot['status'] == 'current' and clans is not None else None
        current_tag = current.get('current_clan_tag') if current else None
        return {**saved_player(row), 'current_clan_tag': current_tag,
                'current_clan_name': current.get('current_clan_name') if current else None,
                'current_clan_type': (clans or {}).get(current_tag, {}).get('type'),
                'current_is_family': current_tag in clans if current else None,
                'current_status': 'current' if current else 'unavailable',
                'current_checked_at': snapshot['checked_at']}

    import asyncio
    players = await asyncio.gather(*(locate(row) for row in found['players']))
    return web.json_response({'players': players, 'total': found['total'], 'limit': limit, 'offset': offset})


def validate_changes(changes):
    """Validate request shape; the storage owner validates changed numeric values."""
    if not isinstance(changes, list) or not 1 <= len(changes) <= 50:
        raise ValueError('Choose between 1 and 50 weights to save.')
    seen = set()
    cleaned = []
    for item in changes:
        if not isinstance(item, dict) or set(item) not in ({'tag', 'weight', 'revision'}, {'tag', 'weight', 'revision', 'storages'}):
            raise ValueError('Invalid weight row')
        player = tag(item['tag'])
        if player in seen or type(item['revision']) is not int or not 0 <= item['revision'] < 2**63:
            raise ValueError('Duplicate player or invalid weight revision')
        seen.add(player)
        if 'storages' in item and (type(item['storages']) is not int or not 1 <= item['storages'] <= 5):
            raise ValueError('Storages must be a whole number from 1 to 5.')
        if item['weight'] is not None and type(item['weight']) is not int:
            raise ValueError('Full weight must be a whole multiple of 1,000.')
        cleaned.append({**item, 'tag': player})
    return cleaned


async def save_player(request):
    from .app import _json, _fail
    actor = authorize(request)
    body = await _json(request)
    try:
        player = tag(request.match_info['tag'])
        if not isinstance(body, dict) or set(body) not in ({'weight', 'revision'}, {'weight', 'revision', 'storages'}):
            raise ValueError('Invalid saved-player weight row')
        change = validate_changes([{**body, 'tag': player}])[0]
        row = request.app['clan_weight_trial'].save_existing(
            player, change, actor['user_id'], actor.get('display_name') or actor.get('username'))
    except WeightConflict as exc:
        _fail(409, str(exc))
    except (ValueError, TypeError) as exc:
        _fail(400, str(exc))
    return web.json_response({'player': saved_player(row)})


async def members(request):
    authorize(request)
    roster = await current(request)
    saved = request.app['clan_weight_trial'].for_members([row['tag'] for row in roster['members']])
    result = response(roster,saved)
    result['last_saved'] = last_saved(request, roster['tag'])
    try:
        lineup = await request.app['providers'].weight_war_lineup(roster['tag'])
        request.app['clan_weight_trial'].remember_lineup(roster['tag'], lineup)
        result['war_order'] = lineup['order']
        result['war_source'] = {key: value for key, value in lineup.items() if key != 'order'}
    except (ProviderUnavailable, ValueError, TypeError):
        lineup = request.app['clan_weight_trial'].saved_lineup(roster['tag'])
        result['war_order'] = lineup['order'] if lineup else None
        result['war_source'] = {key: value for key, value in lineup.items() if key != 'order'} if lineup else None
    return web.json_response(result)


async def save(request):
    from .app import _json, _fail
    actor = authorize(request)
    body = await _json(request)
    try:
        changes = validate_changes(body.get('changes'))
    except (ValueError, TypeError) as exc:
        _fail(400, str(exc))
    seen = {item['tag'] for item in changes}
    roster = await current(request)  # Recheck membership before any write.
    if not seen <= {row['tag'] for row in roster['members']}:
        _fail(409, 'A selected member has left this clan. Refresh members; their previously saved weight is retained.')
    try:
        saved = request.app['clan_weight_trial'].save(roster,changes,actor['user_id'],actor.get('display_name') or actor.get('username'))
    except WeightConflict as error:
        _fail(409,str(error))
    except ValueError as error:
        _fail(400, str(error))
    result = response(roster,saved)
    result['last_saved'] = last_saved(request, roster['tag'])
    return web.json_response(result)


def workbook(rows):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    book = Workbook()
    sheet = book.active
    sheet.title = 'Clan weights'
    headers = ['Player tag','Player name','Town Hall','Weight (k)','Weight','Clan at last save','Clan tag at last save','Saved at','Storages','Loot']
    sheet.append(headers)
    for row in rows:
        values = [row['tag'],row['player_name'],row['town_hall'],row['weight']/1000 if row['weight'] else None,
                  row['weight'],row['last_clan_name'],row['last_clan_tag'],row['updated_at'], row.get('storages'), row['weight'] // row['storages'] if row['weight'] is not None and row.get('storages') else None]
        values = [ILLEGAL_CHARACTERS_RE.sub('', value) if isinstance(value,str) else value for value in values]
        sheet.append(values)
        # Never interpret external names as spreadsheet formulas.
        for cell,value in zip(sheet[sheet.max_row],values):
            if isinstance(value,str): cell.data_type = 's'
    sheet.freeze_panes='A2'
    sheet.auto_filter.ref=sheet.dimensions
    for cell in sheet[1]:
        cell.font=Font(bold=True,color='FFFFFF')
        cell.fill=PatternFill('solid',fgColor='292929')
    for letter,width in zip('ABCDEFGHIJ',[18,30,12,14,16,28,24,30,12,14]):
        sheet.column_dimensions[letter].width=width
    data=BytesIO(); book.save(data)
    return data.getvalue()


async def export(request):
    authorize(request)
    data = workbook(request.app['clan_weight_trial'].all())
    return web.Response(body=data, content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                        headers={'Content-Disposition':'attachment; filename="clan-weights.xlsx"'})


routes = [web.get('/api/cwl/weight-trial/clans',directory),
          web.get('/api/cwl/weight-trial/saved-players', saved_players),
          web.put('/api/cwl/weight-trial/saved-players/{tag}', save_player),
          web.get('/api/cwl/weight-trial/export.xlsx',export),
          web.get('/api/cwl/weight-trial/clans/{tag}',members),
          web.put('/api/cwl/weight-trial/clans/{tag}',save)]
